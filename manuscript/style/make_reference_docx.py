"""Build the Word style template for JCRP submissions (manuscript/style/reference.docx).

Starts from pandoc's default reference.docx and applies manuscript/journal.yaml's format
rules that a style can carry: double spacing, 2.5 cm margins, page numbers at the bottom,
Times New Roman 12 pt, headings in title case as typed (no ALL CAPS, no theme colours).

Usage: uv run python manuscript/style/make_reference_docx.py <pandoc-default.docx> <out.docx>
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
BODY_STYLES = ["Normal", "Body Text", "First Paragraph", "Compact", "Abstract", "Bibliography",
               "Block Text", "Footnote Text", "Caption", "Table Caption", "Image Caption", "Title",
               "Subtitle", "Author", "Date"]
HEADING_STYLES = ["Heading 1", "Heading 2", "Heading 3", "Abstract Title"]


def set_font(style, size, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.all_caps = False
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)


def double_space(style):
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def main(src, out):
    doc = Document(src)
    styles = {s.name: s for s in doc.styles}
    for name in BODY_STYLES:
        if name in styles:
            set_font(styles[name], 12)
            double_space(styles[name])
    for name in HEADING_STYLES:
        if name in styles:
            set_font(styles[name], 12, bold=True)
            double_space(styles[name])
            styles[name].paragraph_format.space_before = Pt(12)
    for section in doc.sections:
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, side, Cm(2.5))
        add_page_number_footer(section)
    doc.save(out)


if __name__ == "__main__":
    main(*sys.argv[1:3])
