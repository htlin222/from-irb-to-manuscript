"""End-to-end tests: config → selected forms → DOCX files."""
import glob
import os

import pytest
from docx import Document

from scripts import generate_all
from scripts.config import load_config
from scripts.form_selector import FORM_REGISTRY, PHASE_FORMS, get_generator, select_forms

FIXTURES = ["config.toml"] + sorted(glob.glob("examples/*/config.toml"))


@pytest.fixture
def retro_config():
    return load_config("examples/gcsf-retrospective/config.toml")


@pytest.fixture
def output_dir(tmp_path):
    path = tmp_path / "output"
    path.mkdir()
    return str(path)


def docx_text(path):
    """All paragraph and table-cell text of a DOCX."""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def generate_phase(config, phase, output_dir):
    config["phase"] = phase
    paths = []
    for fid, _ in select_forms(config):
        path = generate_all.generate_form(fid, config, output_dir)
        assert os.path.getsize(path) > 0, f"{fid} file is empty: {path}"
        paths.append(path)
    return paths


@pytest.mark.parametrize("form_id", sorted(FORM_REGISTRY))
def test_every_registered_generator_runs(form_id, retro_config, output_dir):
    """Each registry entry points at an importable function that writes a DOCX."""
    assert get_generator(form_id) is not None
    path = generate_all.generate_form(form_id, retro_config, output_dir)
    assert os.path.exists(path)


@pytest.mark.parametrize("phase", sorted(PHASE_FORMS))
@pytest.mark.parametrize("fixture", FIXTURES, ids=lambda p: os.path.dirname(p) or "root")
def test_generate_all_every_fixture_and_phase(fixture, phase, tmp_path):
    """The orchestrator succeeds for every shipped fixture in every phase."""
    out = tmp_path / "output"
    code = generate_all.main(fixture, str(out), phase=phase, checklist_path=str(tmp_path / "checklist.md"))
    assert code == 0
    assert len(list(out.glob("*.docx"))) == len(select_forms(load_config(fixture) | {"phase": phase}))


def test_new_case_generates_all_forms(retro_config, output_dir):
    paths = generate_phase(retro_config, "new", output_dir)
    assert len(paths) == 6  # SF001, SF002, SF094, PROPOSAL, SF003, SF005


def test_closure_generates_all_forms(retro_config, output_dir):
    paths = generate_phase(retro_config, "closure", output_dir)
    assert len(paths) == 4  # SF036, SF037, SF038, SF023


def test_generate_all_reports_bad_config(tmp_path, capsys):
    bad = tmp_path / "config.toml"
    bad.write_text('phase = "new"\n[study]\n[pi]\n', encoding="utf-8")
    assert generate_all.main(str(bad), str(tmp_path / "out")) == 2
    assert "study.title_zh" in capsys.readouterr().out
    assert generate_all.main(str(tmp_path / "missing.toml"), str(tmp_path / "out")) == 2


def test_docx_contains_irb_number(retro_config, output_dir):
    path = generate_all.generate_form("SF001", retro_config, output_dir)
    assert "20250801A" in docx_text(path), "IRB number not found in SF001"


def test_docx_contains_pi_name(retro_config, output_dir):
    path = generate_all.generate_form("SF002", retro_config, output_dir)
    assert "林協霆" in docx_text(path), "PI name not found in SF002"


def test_checklist_generation(retro_config, tmp_path):
    """Verify checklist.md is generated correctly."""
    from scripts.checklist import generate_checklist

    results = [
        ("SF001", "新案審查送審資料表", "/fake/path.docx", "generated"),
        ("SF002", "研究計畫申請書", "/fake/path2.docx", "generated"),
        ("SF094", "顯著財務利益申報表", None, "error"),
    ]
    checklist_path = str(tmp_path / "checklist.md")
    generate_checklist(retro_config, results, "新案審查", checklist_path)

    with open(checklist_path) as f:
        content = f.read()

    assert "20250801A" in content
    assert "■ SF001" in content  # generated
    assert "■ SF002" in content  # generated
    assert "□ SF094" in content  # error
    assert "irb@kfsyscc.org" in content


def test_config_validation():
    """Verify config loads without error."""
    config = load_config("examples/gcsf-retrospective/config.toml")
    assert config["study"]["irb_no"] == "20250801A"
    assert config["pi"]["name"] == "林協霆"
    assert config["subjects"]["consent_waiver"] is True
    assert len(config["co_pi"]) == 1
    assert config["co_pi"][0]["name"] == "邱倫維"


def test_proposal_summary_uses_config_text(output_dir):
    """proposal.* text fills 中文計畫摘要; absent keys keep the placeholder."""
    from scripts.generators.proposal import generate_proposal_summary

    config = load_config("examples/tdxd-her2low/config.toml")
    text = "\n".join(p.text for p in Document(
        generate_proposal_summary(config, output_dir)).paragraphs)
    assert "DESTINY-Breast04" in text
    assert "1. 主要目的" in text
    assert "（請列出）" not in text
    assert "標準。疾病" in text  # wrapped Markdown lines joined without spaces

    config.pop("proposal")
    text = "\n".join(p.text for p in Document(
        generate_proposal_summary(config, output_dir)).paragraphs)
    assert "納入條件：（請列出）" in text


def test_re_review_fills_responses_from_markdown(output_dir):
    """re_review.responses (## 意見 N) fill SF019 row by row, with {config} placeholders."""
    from scripts.config import parse_markdown
    from scripts.generators.re_review import generate_sf019

    config = load_config("examples/gcsf-retrospective/config.toml")
    config["closure"]["data_safety"]["retention_years"] = 7
    config["re_review"] = {
        "original_phase": "new",
        "responses": parse_markdown(
            "## 意見 1\n\n**審查意見：** 請說明保存年限。\n\n"
            "**回覆：** 保存{closure.data_safety.retention_years}年。\n"),
        "revised_documents": ["中文計畫摘要"],
    }
    path = generate_sf019(config, output_dir)
    reply_table = Document(path).tables[-1]
    assert [c.text for c in reply_table.rows[1].cells] == ["1", "請說明保存年限。", "保存7年。"]
    assert len(reply_table.rows) == 1 + 5  # padded to the form's 5 rows
    text = docx_text(path)
    assert "■ 新案" in text
    assert "■ 其他（請說明）：中文計畫摘要" in text


def test_proposal_appends_data_protection_prose(output_dir):
    from scripts.generators.proposal import generate_proposal_summary

    config = load_config("examples/tdxd-her2low/config.toml")
    config["proposal"]["data_protection"] = "對照表由計畫主持人保管。"
    text = "\n".join(p.text for p in Document(
        generate_proposal_summary(config, output_dir)).paragraphs)
    assert "資料保存期限" in text and "對照表由計畫主持人保管。" in text


def test_amendment_changes_compare_sections_and_mark_edits(output_dir):
    """SF016 lists only changed proposal sections; deleted text is struck, added text underlined."""
    from scripts.config import parse_markdown
    from scripts.generators import amendment

    old = "## 研究設計\n\n以IPTW平衡。\n\n## 附件\n\n- 資料收集表\n"
    new = parse_markdown("## 研究設計\n\n以重疊加權平衡。\n\n## 附件\n\n- 資料收集表\n")
    rows = amendment.proposal_changes(old, new)
    assert rows == [("中文計畫摘要：研究設計", "以IPTW平衡。", "以重疊加權平衡。")]

    config = load_config("examples/gcsf-retrospective/config.toml")
    config["proposal"] = new
    config["amendment"].update(number=1, reasons=["研究設計變更", "統計方法"], baseline="unused")
    amendment.baseline_text = lambda spec: old   # no git needed in the test
    sf016 = Document(amendment.generate_sf016(config, output_dir)).tables[-1]
    before, after = sf016.rows[1].cells[2].paragraphs[0].runs, sf016.rows[1].cells[3].paragraphs[0].runs
    assert [r.text for r in before if r.font.strike] == ["IPTW"]
    assert [r.text for r in after if r.font.underline] == ["重疊加權"]
    assert len(sf016.rows) == 2   # header + one changed section

    text = docx_text(amendment.generate_sf015(config, output_dir))
    assert "第1次修正" in text
    assert "■ 研究設計變更" in text and "■ 其他（請說明）：統計方法" in text
